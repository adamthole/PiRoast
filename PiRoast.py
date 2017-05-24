# PiRoast.py
# Author: Adam Thole

# THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
# IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
# FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
# AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
# LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
# OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN
# THE SOFTWARE.

import sys
import time
import datetime
from flask import Flask
from flask import jsonify, request, render_template
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
from Roaster import Roaster
from Roast import Roast
from Recipe import Recipe
from RecipeSegment import RecipeSegment

app = Flask(__name__)
roaster = Roaster()
roast = Roast()
port = '/dev/ttyUSB0'

def SaveXlsxFile():
    wb = Workbook()
    wsSummary = wb.active
    wsSummary.title = 'Summary'

    wsSummary.append(["Name", roast.name])
    wsSummary.append(["Start Weight", roast.preWeight])
    wsSummary.append(["End Weight", roast.postWeight])
    wsSummary.append(["Weight Loss", roast.getWeightLoss()])
    wsSummary.append(["First Crack", roast.firstCrack])
    wsSummary.append(["Second Crack", roast.secondCrack])

    wsSummary['B5'].number_format = 'mm:ss.00'
    wsSummary['B6'].number_format = 'mm:ss.00'
    wsSummary.column_dimensions['A'].width = 12
    wsSummary.column_dimensions['B'].width = 25

    wsData = wb.create_sheet(title="Data")
    wsData.append(['Time', 'Internal Temp', 'Target Temp', 'Thermocouple Temp', 'Delta Temp', 'Fan Speed', 'Heat', 'PID Output', 'P', 'I', 'D', 'Heat State Index'])
    for log in roaster.log:
        wsData.append(log)

    for index in range(2,len(roaster.log)+2):
        wsData['A' + str(index)].number_format = 'mm:ss.00'

    chart = LineChart()
    chart.title = "Temperature Profile"
    chart.style = 2
    chart.y_axis.title = 'Temperature (F)'
    chart.x_axis.title = 'Time'

    data = Reference(wsData, min_col=3, min_row=1, max_col=4, max_row=len(roaster.log) + 1)
    timings = Reference(wsData, min_col=1, min_row=1, max_col=1, max_row=len(roaster.log) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(timings)

    wsSummary.add_chart(chart, "A9")

    chart2 = LineChart()
    chart2.title = "Roaster Settings"
    chart2.style = 2
    chart2.y_axis.title = 'Setting'
    chart2.x_axis.title = 'Time'

    fanHeatData = Reference(wsData, min_col=5, min_row=1, max_col=6, max_row=len(roaster.log) + 1)
    chart2.add_data(fanHeatData, titles_from_data=True)
    chart2.set_categories(timings)

    wsSummary.add_chart(chart2, "A24")
    # Save the file
    wb.save('roasts/' + roast.startTime.strftime('%y-%m-%d ') + roast.name + ".xlsx")

@app.route("/")
def main():
    return render_template('main.html')

@app.route("/getRecipes")
def getRecipes():
    return jsonify.dumps([1,2,3])
    return roaster.Recipes[0].name;

@app.route('/getRoastInfo')
def getRoastInfo():
    internalTemp = roaster.getInternalTemp()
    tcTemp = roaster.getThermocoupleTemp()
    heatSetting = roaster.getHeatSetting()
    fanSpeed = roaster.getFanSpeed()
    elapsedTime = formatElapsedTime(roaster.getElapsedTime())
    deltaTemp = roaster.getDeltaTemp()
    return jsonify(tcTemp=tcTemp, internalTemp=str(roaster.targetTemp),
            targetTemp = roaster.targetTemp,
            heatSetting=heatSetting,
            fanSpeed = fanSpeed, elapsedTime = elapsedTime,
            firstCrack = formatElapsedTime(roast.firstCrack),
            secondCrack = formatElapsedTime(roast.secondCrack),
            deltaTemp = deltaTemp)

@app.route('/autoTemp/<targetTemp>')
def autoTemp(targetTemp):
    roaster.autoTemp(int(targetTemp))
    return 'OK'

@app.route('/setHeat/<heatSetting>')
def setHeat(heatSetting):
    roaster.setHeat(int(heatSetting))
    return 'OK'

@app.route('/setFanSpeed/<fanSpeed>')
def setFanSpeed(fanSpeed):
    roaster.setFanSpeed(int(fanSpeed))
    return 'OK'

@app.route('/firstCrack')
def firstCrack():
    roast.firstCrack = roaster.getElapsedTime()
    return 'OK'

@app.route('/secondCrack')
def secondCrack():
    roast.secondCrack = roaster.getElapsedTime()
    return 'OK'

@app.route('/startRoast')
def startRoast():
    print('Starting Roast!', file=sys.stderr)
    roast.startTime = datetime.datetime.now()
    roaster.startRoast()
    return str(roast.startTime)

@app.route('/coolRoast')
def coolRoast():
    roaster.coolRoast()
    return 'OK'

@app.route('/stopRoast')
def stopRoast():
    roaster.stopRoast()
    return 'OK'

@app.route('/saveRoast/<roastName>/<preWeight>/<postWeight>')
def saveRoast(roastName, preWeight, postWeight):
    roast.name = roastName
    roast.preWeight = int(preWeight)
    roast.postWeight = int(postWeight)
    SaveXlsxFile()
    return 'OK'

@app.route('/recipe')
def DoMyRecipe():
    roaster.runRecipe()

def formatElapsedTime(elapsedTime):
    s = elapsedTime.seconds
    hours, remainder = divmod(s, 3600)
    minutes, seconds = divmod(remainder, 60)

    if len(str(seconds)) == 1: seconds = "0" + str(seconds)
    return str('%s:%s' % (minutes, seconds));

if __name__ == '__main__':
    roaster.connect(port)
    app.run(host='0.0.0.0', port=80, threaded=True)

    isRun = True

